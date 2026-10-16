import io

import pytest
from docx import Document

from app.documents import extract_document_text
from tests.samples import AI_TEXT


def _docx(text: str) -> bytes:
    doc = Document()
    for para in text.split(". "):
        doc.add_paragraph(para)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf(text: str) -> bytes:
    """A minimal single-page PDF with one line of Helvetica text."""
    stream = f"BT /F1 10 Tf 20 700 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
    ]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objs) + 1, xref)
    )
    return out.getvalue()


def _docx_with_table(
    table_data: list[list[str]], paragraphs: list[str] | None = None
) -> bytes:
    doc = Document()
    if paragraphs:
        for p in paragraphs:
            doc.add_paragraph(p)
    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, cell_value in enumerate(row):
                table.rows[r_idx].cells[c_idx].text = cell_value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx():
    text = extract_document_text("essay.DOCX", _docx(AI_TEXT))
    assert "multifaceted tapestry" in text


def test_docx_text_in_table_only():
    data = [
        ["Project Title", "SlopTotal AI Detection Engine"],
        [
            "Candidate Background",
            "Senior Machine Learning Engineer with 8 years of experience building NLP systems.",
        ],
    ]
    raw = _docx_with_table(data)
    extracted = extract_document_text("resume.docx", raw)
    assert "SlopTotal AI Detection Engine" in extracted
    assert "Senior Machine Learning Engineer" in extracted


def test_docx_paragraphs_then_tables_joined_by_blank_lines():
    paragraphs = ["First paragraph introduction.", "Second paragraph description."]
    table_data = [["Cell One", "Cell Two"]]
    raw = _docx_with_table(table_data, paragraphs=paragraphs)
    extracted = extract_document_text("mixed.docx", raw)
    expected = (
        "First paragraph introduction.\n\n"
        "Second paragraph description.\n\n"
        "Cell One\n\n"
        "Cell Two"
    )
    assert extracted == expected


def test_pdf():
    line = (
        "Missed the bus again this morning, so I walked for forty minutes in the rain."
    )
    assert "forty minutes" in extract_document_text("note.pdf", _pdf(line))


def test_plain_text():
    assert extract_document_text("a.txt", AI_TEXT.encode()) == AI_TEXT


@pytest.mark.parametrize(
    "name,data,msg",
    [
        ("a.exe", b"MZ" * 100, "Supported files"),
        ("a.pdf", b"not a pdf at all", "Could not read"),
        ("a.txt", b"too short", "No readable text"),
        ("a.txt", b"x" * (10 * 1024 * 1024 + 1), "larger than 10 MB"),
    ],
)
def test_rejections(name, data, msg):
    with pytest.raises(ValueError, match=msg):
        extract_document_text(name, data)


@pytest.mark.parametrize("name", ["broken.pdf", "broken.docx"])
def test_parse_errors_preserve_the_original_cause(name):
    with pytest.raises(ValueError, match="Could not read") as caught:
        extract_document_text(name, b"invalid document")
    assert caught.value.__cause__ is not None


def test_markdown_is_decoded_and_returned_as_text():
    # Long enough to pass the 50-character minimum. Markdown syntax is returned as-is, not rendered.
    md = "# Notes\n\n- first point\n- second point\n\nThis paragraph is long enough to pass the minimum."
    assert extract_document_text("notes.md", md.encode("utf-8")) == md


def test_pdf_without_extractable_text_is_rejected():
    # Must be a valid PDF that parses but has no text; match= proves it was rejected for that
    # reason and not because the file could not be read.
    with pytest.raises(ValueError, match="No readable text"):
        extract_document_text("blank.pdf", _pdf(""))


def test_docx_preserves_interleaved_paragraph_and_table_order():
    doc = Document()
    doc.add_paragraph("Introduction before the first table.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "First table contents."
    doc.add_paragraph("Explanation between the two tables.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Second table contents."
    doc.add_paragraph("Conclusion after the second table.")
    buf = io.BytesIO()
    doc.save(buf)

    assert extract_document_text("interleaved.docx", buf.getvalue()) == "\n\n".join(
        [
            "Introduction before the first table.",
            "First table contents.",
            "Explanation between the two tables.",
            "Second table contents.",
            "Conclusion after the second table.",
        ]
    )


def test_docx_preserves_a_table_before_the_first_paragraph():
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Opening table with context."
    doc.add_paragraph("Following paragraph with the explanation of that table.")
    buf = io.BytesIO()
    doc.save(buf)

    assert extract_document_text("table-first.docx", buf.getvalue()) == (
        "Opening table with context.\n\n"
        "Following paragraph with the explanation of that table."
    )


@pytest.mark.parametrize(
    ("merge_end", "expected_cells"),
    [
        (
            (0, 1),
            ["Merged cell contents.", "Lower left contents.", "Lower right contents."],
        ),
        (
            (1, 0),
            ["Merged cell contents.", "Upper right contents.", "Lower right contents."],
        ),
    ],
    ids=["horizontal-merge", "vertical-merge"],
)
def test_docx_extracts_merged_cells_once_in_document_order(merge_end, expected_cells):
    doc = Document()
    doc.add_paragraph("Paragraph before the merged table.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 1).text = "Upper right contents."
    table.cell(1, 0).text = "Lower left contents."
    table.cell(1, 1).text = "Lower right contents."
    table.cell(0, 0).merge(table.cell(*merge_end)).text = "Merged cell contents."
    doc.add_paragraph("Paragraph after the merged table.")
    buf = io.BytesIO()
    doc.save(buf)

    assert extract_document_text("merged.docx", buf.getvalue()) == "\n\n".join(
        [
            "Paragraph before the merged table.",
            *expected_cells,
            "Paragraph after the merged table.",
        ]
    )
